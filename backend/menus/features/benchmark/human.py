# AI-generated: written with Claude (Anthropic) and reviewed by the team.
"""Human reference ratings: raters score dishes with the extractor's rubric.

The form (export_feature_rating_form) is one file per rater, an Excel workbook or a CSV:

    name, kind, rater, recipe_notes, spicy_mean, spicy_std, sweet_mean, sweet_std, ...

The workbook's sheet "평가" holds these columns from the row whose first cell is
`name`, with the rater's name in B1 instead of a `rater` column; its other sheets are
the rater's guide and the rubric.

- `<key>_mean`, `<key>_std`: 0 to 1, at most two decimals, std² <= mean * (1 - mean);
  the same rules as the extractor's answer (menus/features/instructions.txt)
- `kind`: what kind of dish it is (고정 메뉴, 매일 바뀌는 메뉴, ...), for the rater only
- `recipe_notes`: the recipes or serving examples looked at before judging the stds
- a blank cell is skipped; a header may carry a label after a colon ("spicy_mean: 매운맛 평균")
- without a `rater` column, the file name is the rater

Comparison, on the dishes and features both raters (or the LLM and a rater) scored:
- rater agreement: mean absolute difference between two raters, averaged over rater pairs
- LLM vs raters: mean absolute difference between the LLM and each rater, averaged over raters.
  Same scale, so a ratio near 1 means the LLM is as far from people as people are from each other.
- Spearman correlation of the LLM's means with the raters' average
Means and stds are compared separately.
"""

from __future__ import annotations

import csv
import math
import pathlib
from collections import defaultdict
from itertools import combinations

import numpy as np
import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from menus.features.schema import LLM_FEATURE_KEYS, RUBRIC

FIELDS = ("mean", "std")
SHEET = "평가"


def _number(text, where):
    try:
        value = float(text)
    except ValueError:
        raise ValueError(f"{where}: {text!r} is not a number") from None
    if not math.isfinite(value) or abs(value - round(value, 2)) > 1e-9:
        raise ValueError(f"{where}: {text!r} must be a number with at most two decimals")
    return value


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _rows(path):
    """(where, {column: text}) for each data row of a CSV or Excel form."""
    if path.suffix != ".xlsx":
        with open(path, encoding="utf-8-sig", newline="") as f:
            for line, row in enumerate(csv.DictReader(f), start=2):
                yield f"line {line}", {(key or "").split(":")[0].strip(): _text(value) for key, value in row.items()}
        return
    book = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in book.sheetnames:
        raise ValueError(f"{path.name}: no sheet {SHEET!r}")
    rows = list(book[SHEET].iter_rows(values_only=True))
    start = next((i for i, row in enumerate(rows) if row[0] == "name"), None)
    if start is None:
        raise ValueError(f"{path.name}: no header row starting with 'name' in sheet {SHEET!r}")
    rater, header = _text(rows[0][1]), [_text(value) for value in rows[start]]
    for number, row in enumerate(rows[start + 1:], start=start + 2):
        yield f"row {number}", {"rater": rater, **{c: _text(v) for c, v in zip(header, row) if c}}


def load(paths) -> dict[str, dict[str, dict]]:
    """{rater: {dish: {"mean": {key: value}, "std": {key: value}, "notes": str}}}."""
    ratings = defaultdict(dict)
    for path in map(pathlib.Path, paths):
        for where, row in _rows(path):
            name, rater = row.pop("name", ""), row.pop("rater", "") or path.stem
            notes = row.pop("recipe_notes", "")
            row.pop("kind", None)
            if not name:
                continue
            entry = {"mean": {}, "std": {}, "notes": notes}
            for column, text in row.items():
                key, _, field = column.rpartition("_")
                if key not in LLM_FEATURE_KEYS or field not in FIELDS:
                    raise ValueError(f"{path.name}: unknown column {column!r}")
                if text:
                    entry[field][key] = _number(text, f"{path.name} {where}, {column}")
            for field, top in (("mean", 1), ("std", 0.5)):
                for key, value in entry[field].items():
                    if not 0 <= value <= top:
                        raise ValueError(f"{path.name} {where}, {key}_{field}: {value} is outside 0..{top}")
            for key in entry["mean"].keys() & entry["std"].keys():
                mean, std = entry["mean"][key], entry["std"][key]
                if std * std > mean * (1 - mean) + 1e-9:
                    raise ValueError(f"{path.name} {where}, {key}: mean {mean} / std {std} breaks "
                                     "std² <= mean(1 - mean)")
            ratings[rater][name] = entry
    return dict(ratings)


def reference(ratings) -> dict[str, tuple[dict, dict]]:
    """{dish: (means, stds)} averaged over raters, for dishes with every feature rated."""
    result = {}
    for name in sorted({name for dishes in ratings.values() for name in dishes}):
        entries = [dishes[name] for dishes in ratings.values() if name in dishes]
        averaged = []
        for field in FIELDS:
            values = {key: [e[field][key] for e in entries if key in e[field]] for key in LLM_FEATURE_KEYS}
            if not all(values.values()):
                break
            averaged.append({key: float(np.mean(v)) for key, v in values.items()})
        else:
            result[name] = tuple(averaged)
    return result


def spearman(a, b) -> float:
    def ranks(values):
        values = np.asarray(values, dtype=float)
        result = np.empty(len(values))
        result[values.argsort(kind="stable")] = np.arange(len(values))
        for value in np.unique(values):  # ties share their average rank
            result[values == value] = result[values == value].mean()
        return result
    if len(a) < 3 or len(set(a)) < 2 or len(set(b)) < 2:
        return float("nan")
    return float(np.corrcoef(ranks(a), ranks(b))[0, 1])


def _mean(values) -> float:
    return float(np.mean(values)) if values else float("nan")


def compare(runset, ratings, field) -> tuple[list, dict]:
    """Per-feature rows and overall numbers for `field` ("mean" or "std"), over the rated dishes
    in `runset`: the raters' agreement is measured on the same dishes as the LLM's."""
    raters = sorted(ratings)
    index = {key: i for i, key in enumerate(runset.keys)}
    llm = runset.mean if field == "mean" else runset.std
    names = [n for n in sorted({n for r in raters for n in ratings[r]}) if llm(n) is not None]
    rows, pooled = [], defaultdict(list)
    for key in LLM_FEATURE_KEYS:
        if key not in index:
            continue
        between = [abs(ratings[a][n][field][key] - ratings[b][n][field][key]) for a, b in combinations(raters, 2)
                   for n in names if n in ratings[a] and n in ratings[b]
                   and key in ratings[a][n][field] and key in ratings[b][n][field]]
        against, model, people = [], [], []
        for name in names:
            values = [ratings[r][name][field][key] for r in raters if key in ratings[r].get(name, {}).get(field, {})]
            vector = llm(name)
            if values:
                against += [abs(vector[index[key]] - v) for v in values]
                model.append(float(vector[index[key]]))
                people.append(float(np.mean(values)))
        pooled["between"] += between
        pooled["against"] += against
        rater_gap, llm_gap = _mean(between), _mean(against)
        rows.append([key, len(model), round(rater_gap, 3), round(llm_gap, 3),
                     round(llm_gap / rater_gap, 2) if rater_gap > 0 else "", round(spearman(model, people), 2)])
    overall = {"raters": _mean(pooled["between"]), "llm": _mean(pooled["against"])}
    return rows, overall


def evaluate(runset, ratings) -> dict:
    title = "4. Human reference"
    if not ratings:
        return {"title": title, "skipped": "No rating file. Make the form with export_feature_rating_form "
                "and pass the filled copies with --ratings."}
    rated = {name for dishes in ratings.values() for name in dishes}
    shared = sorted(n for n in rated if runset.mean(n) is not None)
    if not shared:
        return {"title": title, "skipped": "None of the rated dishes are in the runs file."}
    lines = [f"{len(shared)} dishes rated by {len(ratings)} rater(s): {', '.join(sorted(ratings))}.",
             "ratio: how far the LLM is from each rater over how far raters are from each other. About 1: the LLM "
             "looks like one more rater. An LLM at the raters' true average would score about 0.7 (1/√2), "
             "since each rater's own error remains."]
    tables = []
    for field in FIELDS:
        rows, overall = compare(runset, ratings, field)
        ratio = overall["llm"] / overall["raters"] if overall["raters"] > 0 else float("nan")
        lines.append(f"{field}: LLM vs raters {overall['llm']:.3f}, rater vs rater "
                     + (f"{overall['raters']:.3f} (ratio {ratio:.2f})" if not math.isnan(overall["raters"])
                        else "- (one rater: no agreement to compare with)"))
        tables.append((f"human_{field}.csv", f"Per feature, {field}s",
                       ["feature", "dishes", "rater vs rater", "LLM vs raters", "ratio", "Spearman"], rows, len(rows)))
    return {"title": title, "lines": lines, "tables": tables}


def write_form(path, dishes, values=None, rater="", note="") -> None:
    """A rating form, Excel for a .xlsx path and CSV otherwise: one row per (name, kind)
    in `dishes`, a mean and a std column per rubric feature. Blank for raters, or
    filled by `rater` with `values` ({name: (means, stds)}), e.g. the LLM's answers."""
    if pathlib.Path(path).suffix == ".xlsx":
        _write_workbook(path, dishes, values, rater, note)
        return
    labels = {feature["key"]: feature["label_ko"] for feature in RUBRIC["features"]}
    header = ["name", "kind", "rater", "recipe_notes"]
    for key in LLM_FEATURE_KEYS:
        header += [f"{key}_mean: {labels[key]} 평균", f"{key}_std: {labels[key]} 표준편차"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for name, kind in dishes:
            writer.writerow([name, kind, rater, ""] + _values(values, name))


FONT, BOLD, KEY_FONT = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True), \
    Font(name="Arial", size=8, color="808080")
INPUT, FIXED, HEAD = (PatternFill("solid", fgColor=c) for c in ("FFF2CC", "EDEDED", "DDEBF7"))
BROKEN = PatternFill("solid", fgColor="F4B6B6")
WRAP = Alignment(wrap_text=True, vertical="top")
# The extractor's instructions up to its output format, which is only for the LLM.
GUIDE = (pathlib.Path(__file__).parents[1] / "instructions.txt").read_text(encoding="utf-8").split("\n출력\n")[0]
EXAMPLE = [("name", "김치찌개 (예시)"), ("kind", "고정 메뉴"), ("recipe_notes", "레시피 3개: 돼지고기·참치 변형, 김치 숙성도 차이"),
           ("spicy_mean", 0.6), ("spicy_std", 0.15), ("sour_mean", 0.4), ("sour_std", 0.15),
           ("pork_mean", 0.6), ("pork_std", 0.35), ("noodle_mean", 0), ("noodle_std", 0)]


def _values(values, name) -> list:
    if not values:
        return [""] * (2 * len(LLM_FEATURE_KEYS))
    means, stds = values[name]
    return [value for key in LLM_FEATURE_KEYS for value in (means[key], stds[key])]


def _cell(sheet, row, column, value, font=FONT, fill=None):
    cell = sheet.cell(row, column, value)
    cell.font, cell.alignment = font, WRAP
    if fill:
        cell.fill = fill
    return cell


def _write_workbook(path, dishes, values, rater, note) -> None:
    book = openpyxl.Workbook()
    form = book.active
    form.title = SHEET
    blank = INPUT if not values else None
    _cell(form, 1, 1, "평가자 이름", BOLD)
    _cell(form, 1, 2, rater or None, fill=blank)
    _cell(form, 1, 3, note or "노란 칸만 입력하세요. 작성 방법은 '안내', 기준은 '루브릭' 시트에 있습니다.")
    form.merge_cells(start_row=1, start_column=3, end_row=1, end_column=12)
    first, last = 5, 4 + len(dishes)
    labels = {feature["key"]: feature["label_ko"] for feature in RUBRIC["features"]}
    columns = ["name", "kind", "recipe_notes"] + [f"{key}_{field}" for key in LLM_FEATURE_KEYS for field in FIELDS]
    for index, text in enumerate(("음식", "구분", "참고한 레시피·제공 예시"), start=1):
        _cell(form, 3, index, text, BOLD, HEAD)
    for index, column in enumerate(columns, start=1):
        _cell(form, 4, index, column, KEY_FONT, HEAD)
    for row, (name, kind) in enumerate(dishes, start=first):
        _cell(form, row, 1, name, fill=FIXED)
        _cell(form, row, 2, kind, fill=FIXED)
        _cell(form, row, 3, None, fill=blank)
        for column, value in enumerate(_values(values, name), start=4):
            _cell(form, row, column, value if values else None, fill=blank)

    ranges = {}
    for field, top in (("mean", 1), ("std", 0.5)):
        ranges[field] = DataValidation(type="decimal", operator="between", formula1="0", formula2=str(top),
                                       allow_blank=True, showErrorMessage=True, errorTitle="범위 밖",
                                       error=f"{field}는 0~{top} 사이 숫자입니다.")
        form.add_data_validation(ranges[field])
    for index, key in enumerate(LLM_FEATURE_KEYS):
        column = 4 + 2 * index
        mean, std = get_column_letter(column), get_column_letter(column + 1)
        _cell(form, 3, column, f"{labels[key]} ({key})", BOLD, HEAD)
        form.merge_cells(start_row=3, start_column=column, end_row=3, end_column=column + 1)
        ranges["mean"].add(f"{mean}{first}:{mean}{last}")
        ranges["std"].add(f"{std}{first}:{std}{last}")
        form.conditional_formatting.add(f"{mean}{first}:{std}{last}", FormulaRule(fill=BROKEN, formula=[
            f"AND(ISNUMBER(${mean}{first}),ISNUMBER(${std}{first}),"
            f"${std}{first}^2>${mean}{first}*(1-${mean}{first})+0.000000001)"]))
    end = get_column_letter(len(columns))
    form.conditional_formatting.add(f"D{first}:{end}{last}", FormulaRule(fill=BROKEN, formula=[
        f"AND(ISNUMBER(D{first}),ABS(D{first}-ROUND(D{first},2))>0.000000001)"]))
    for letter, width in zip("ABC", (20, 14, 28)):
        form.column_dimensions[letter].width = width
    for index in range(4, len(columns) + 1):
        form.column_dimensions[get_column_letter(index)].width = 7
    form.freeze_panes = f"B{first}"
    if not values:
        _write_guide(book.create_sheet("안내"))
    _write_rubric(book.create_sheet("루브릭"))
    book.save(path)


def _write_guide(guide) -> None:
    lines = [
        ("Metchu! feature 평가 양식", BOLD),
        ("음식 이름으로 추출한 32개 feature를 사람의 평가와 비교하기 위한 기준 데이터입니다. "
         "평가자마다 이 파일을 한 부씩 받아 같은 음식을 각자 평가합니다.", FONT),
        ("구분: 고정 메뉴는 학식의 고정 메뉴판 음식, 매일 바뀌는 메뉴는 날마다 바뀌는 식단의 음식, "
         "비교용 예시는 비슷한 음식과 구분되는지 보려고 넣은 음식입니다. 평가 방법은 모두 같습니다.", FONT),
        ("", FONT), ("작성 순서", BOLD),
        ("1. '평가' 시트 B1에 이름을 적습니다.", FONT),
        ("2. 음식마다 레시피와 식당 제공 예시(사진·메뉴 설명)를 2~3개 찾아보고, 참고한 내용을 recipe_notes에 짧게 적습니다. "
         "std는 이렇게 본 레시피·구성 차이를 근거로 정합니다.", FONT),
        ("3. 각 feature의 _mean(평균)과 _std(표준편차)를 '루브릭' 시트와 아래 지침에 따라 적습니다. "
         "다른 평가자나 LLM의 값은 보지 않습니다.", FONT),
        ("4. 빨간 칸이 없는지 확인하고 .xlsx 그대로 저장해 제출합니다. 32개 feature를 모두 채운 음식만 추천 영향 평가에 쓰입니다.", FONT),
        ("", FONT), ("규칙", BOLD),
        ("mean은 0~1, std는 0~0.5, 소수점 둘째 자리까지. std² ≤ mean×(1−mean)이어야 하고, mean이 0 또는 1이면 std는 0입니다.", FONT),
        ("", FONT), ("범례", BOLD),
    ]
    for row, (text, font) in enumerate(lines, start=1):
        _cell(guide, row, 1, text, font)
    row = len(lines) + 1
    for text, fill in (("입력하는 칸", INPUT), ("수정하지 않는 칸 (음식 이름)", FIXED),
                       ("규칙 위반 (범위·소수 자리·std² ≤ mean×(1−mean))", BROKEN)):
        _cell(guide, row, 1, text, fill=fill)
        row += 1
    _cell(guide, row + 1, 1, "작성 예시 (형식 확인용, 일부 feature만 표시)", BOLD)
    for index, (column, value) in enumerate(EXAMPLE, start=1):
        _cell(guide, row + 2, index, column, KEY_FONT, HEAD)
        _cell(guide, row + 3, index, value, fill=INPUT)
    row += 5
    _cell(guide, row, 1, "추출 지침 (LLM과 같은 기준)", BOLD)
    for row, text in enumerate(GUIDE.strip().splitlines(), start=row + 1):
        _cell(guide, row, 1, text)
    guide.column_dimensions["A"].width = 100
    for index in range(2, len(EXAMPLE) + 1):
        guide.column_dimensions[get_column_letter(index)].width = 12


def _write_rubric(rubric) -> None:
    header = ["key", "이름", "정의", "구분 규칙", "평균 기준", "표준편차 기준"]
    for index, text in enumerate(header, start=1):
        _cell(rubric, 1, index, text, BOLD, HEAD)
    for row, feature in enumerate((f for f in RUBRIC["features"] if f["key"] in LLM_FEATURE_KEYS), start=2):
        values = [feature["key"], feature["label_ko"], feature["definition"], "\n".join(feature["distinction_rules"]),
                  feature["mean_guidance"], feature["std_guidance"]]
        for index, value in enumerate(values, start=1):
            _cell(rubric, row, index, value)
    for letter, width in zip("ABCDEF", (14, 10, 40, 60, 50, 50)):
        rubric.column_dimensions[letter].width = width
    rubric.freeze_panes = "C2"
